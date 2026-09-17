"""Build the shareable Neural Pods PowerPoint deck."""
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parents[1] / "deliverables" / "neural-pods-presentation.pptx"
BG = RGBColor(11,16,32); CYAN = RGBColor(139,233,253); GREEN = RGBColor(80,250,123); PURPLE = RGBColor(189,147,249); MUTED = RGBColor(170,180,208)

def add_text(slide, text, x, y, w, h, size=24, color=RGBColor(238,242,255), bold=False):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=box.text_frame; tf.word_wrap=True
    p=tf.paragraphs[0]; p.text=text; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=color; p.font.name='Aptos'; return box

def slide(prs, title, body, accent=CYAN):
    s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=BG
    add_text(s,title,.7,.5,12,0.7,30,accent,True); add_text(s,body,.9,1.45,11.5,5.4,22)
    add_text(s,'Neural Pods · lokaler Forschungsprototyp · 16.09.2026',.7,7.05,12,.25,9,MUTED)
    return s

prs=Presentation(); prs.slide_width= Inches(13.333); prs.slide_height= Inches(7.5)
s=slide(prs,'Neural Pods','Versioniertes internes Wissen für reale Sprachmodelle\n\nOriginKey → KnowledgeKey → GenerationKey\n→ Dragonfly → Pod/LoRA → Qwen3B',CYAN)
slide(prs,'Das Problem mit reinem RAG','Bei jeder Anfrage: Dokumente suchen, Kontext bauen, Präfix wiederholen.\n\nBei Updates: Index, Cache und materialisierte Antworten müssen konsistent widerrufen werden.\n\nFrage → Retrieval → langer Kontext → Modell',PURPLE)
slide(prs,'Die Pod-Idee','Quelle\n  │ OriginKey\n  ▼\nKanonisches Wissen\n  ├─ KnowledgeKey / GenerationKey\n  ├─ Provenienz-DAG\n  ├─ Embedding + BM25-Adresse\n  └─ generation-bound neural payload\n             ▼\n          Qwen3B',CYAN)
slide(prs,'Ein gemeinsamer Lebenszyklus','K:g7 ──► Pod:A91 ──► Cache:C144\n  │\n  └─ revoke\n       ↓\nK:g8 ──► Pod:A92\n\nAliases zeigen auf dieselbe semantische Identität. Revocation propagiert über abgeleitete Artefakte.',PURPLE)
slide(prs,'Was ist implementiert?','Semantik: Pod-Typen, Domain, Rollen, Intent, Entitäten, Alias-Familien\n\nRetrieval: lokale Native Embeddings, ANN/kNN-ähnliche Suche, BM25, Filter, RRF\n\nNeural: Dragonfly-Router und generation-bound Qwen3B-LoRA\n\nLifecycle: Origin, Generation, ACL, Branching, Revocation, Commit-Barrier',CYAN)
slide(prs,'Reales Qwen3B-Messergebnis','Kompakter Pod-Kontext: 61 Eingabetoken · 190,9 ms pro Batch (4)\n\nLängerer RAG-Kontext: 303 Eingabetoken · 412,5 ms pro Batch (4)\n\n→ 2,16× niedrigere Pod-Latenz\n→ Beide Pfade liefern im Fixture exakt „18 days“',GREEN)
slide(prs,'Weitere Messwerte','23,9 Token/s · Qwen3B auf RTX 3090\n9,48 ms p50 · lokale Suche bei 2.000 Rows\n434× · inkrementelles R211b-Update\n20/20 · Projekt-Gate-Checks bestanden',GREEN)
slide(prs,'Was der Nachweis bedeutet','Die Architektur läuft auf einem echten Qwen3B-Modell.\n\nProvenienz und Generation sind ausführbar und widerrufbar.\n\nDer kompakte Pod-Pfad reduziert Prefill-Last im kontrollierten Vergleich.\n\nDie lokale Search-Baseline ist noch kein 100B-ANN-System.',PURPLE)
slide(prs,'Nächster wissenschaftlicher Schritt','100–1.000 mutable facts\nPod-LoRA vs. optimized RAG vs. hybrid\n→ paraphrase · multi-hop · update · revoke\n→ accuracy · latency · VRAM · QPS\n\nDie Hypothese ist testbar: weniger Kontext bei kontrollierbarerem Lebenszyklus.',CYAN)
OUT.parent.mkdir(parents=True,exist_ok=True); prs.save(OUT); print(OUT)
